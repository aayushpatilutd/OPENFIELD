#!/usr/bin/env python3
"""
End-to-end smoke test.

Boots the app with Flask's test client, signs in, and walks the whole AP cycle
the way a person would: sweep, parse, sync, push, reconcile, export. Every page
is rendered and every API route is called, so a template typo or a bad SQL
string fails here rather than in front of someone.

Run:  python tests/smoke.py
"""

import os
import sys
import json
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

# A throwaway database, so running the tests never touches real data.
TMP = tempfile.mkdtemp(prefix="invoicehub-test-")
os.environ["INVOICEHUB_DB"] = os.path.join(TMP, "test.db")
os.environ["INVOICEHUB_STORAGE"] = os.path.join(TMP, "storage")
os.environ["INVOICEHUB_DEMO"] = "1"
os.environ["INVOICEHUB_SECRET_KEY"] = "test-key"

import db                      # noqa: E402
from app import app            # noqa: E402

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  {'PASS' if cond else 'FAIL'}  {name}"
          + (f"   — {detail}" if detail and not cond else ""))
    return cond


def main():
    app.config["TESTING"] = True
    c = app.test_client()

    print("\n── auth ──────────────────────────────────────────────")
    r = c.get("/")
    check("anonymous is redirected to login", r.status_code == 302)
    r = c.get("/api/vendors")
    check("API refuses anonymous with 401", r.status_code == 401)

    r = c.post("/login", data={"email": "admin@invoicehub.local",
                               "password": "wrong"})
    check("a wrong password is rejected", b"do not match" in r.data)

    r = c.post("/login", data={"email": "admin@invoicehub.local",
                               "password": "changeme"},
               follow_redirects=False)
    check("correct password signs in", r.status_code == 302)

    print("\n── demo data ─────────────────────────────────────────")
    r = c.post("/api/demo/seed", json={})
    d = r.get_json()
    check("demo seed runs", d.get("ok"), json.dumps(d)[:300])
    check("vendors were created", d.get("vendors", 0) >= 10)
    check("the seed reconciles", d.get("match", {}).get("ok"))

    print("\n── pages render ──────────────────────────────────────")
    for path in ("/", "/vendors", "/inbox", "/invoices", "/statements",
                 "/reconciliation", "/reports", "/agents", "/drafts",
                 "/settings", "/activity"):
        r = c.get(path)
        ok = r.status_code == 200 and b"<!DOCTYPE html>" in r.data
        check(f"GET {path}", ok, f"status {r.status_code}")

    print("\n── read APIs ─────────────────────────────────────────")
    for path in ("/api/dashboard", "/api/vendors", "/api/documents",
                 "/api/invoices", "/api/statements", "/api/match/report",
                 "/api/match/plan", "/api/reports", "/api/agents",
                 "/api/drafts", "/api/activity", "/api/settings", "/api/me",
                 "/api/users", "/api/scan/status"):
        r = c.get(path)
        j = r.get_json() or {}
        check(f"GET {path}", r.status_code == 200 and j.get("ok") is not False,
              f"status {r.status_code} {str(j)[:160]}")

    print("\n── the AP cycle ──────────────────────────────────────")
    r = c.post("/api/scan", json={})
    d = r.get_json()
    check("mailbox sweep runs", d.get("ok"), str(d)[:200])
    check("the sweep files new documents", d.get("filed", 0) > 0,
          f"filed={d.get('filed')}")
    check("the duplicate attachment is recognised, not refiled",
          d.get("skipped", 0) > 0, f"skipped={d.get('skipped')}")
    check("mail from an unknown sender is held for a human",
          d.get("unmatched", 0) > 0, f"unmatched={d.get('unmatched')}")

    r = c.post("/api/qbo/sync", json={})
    d = r.get_json()
    check("QuickBooks sync runs", d.get("ok"), str(d)[:200])
    check("bills came back", d.get("bills", 0) > 0)

    r = c.post("/api/match/run", json={})
    d = r.get_json()
    check("reconciliation runs", d.get("ok"), str(d)[:200])
    check("rows were produced", d.get("rows", 0) > 0)

    print("\n── the verdicts the demo must demonstrate ────────────")
    rep = c.get("/api/match/report").get_json()
    verdicts = {}
    for g in rep["groups"]:
        for row in g["rows"]:
            verdicts[row["verdict"]] = verdicts.get(row["verdict"], 0) + 1
    print(f"    {verdicts}")
    for v in ("matched", "amount_mismatch", "missing_bill", "credit",
              "statement_only", "qbo_only"):
        check(f"demo covers '{v}'", verdicts.get(v, 0) > 0)

    print("\n── invoice-number normalisation ──────────────────────")
    # ABC-000001 on the PDF, 1 on the statement, INV-000001 in QBO.
    crl = next((g for g in rep["groups"] if g["vendor_name"] == "Company ABC"), None)
    check("Company ABC is in the report", crl is not None)
    if crl:
        row = next((r for r in crl["rows"]
                    if r["invoice_number"] == "ABC-000001"), None)
        check("the three spellings fold into one row", row is not None)
        if row:
            check("and all three amounts are present on it",
                  row["amt_invoice"] and row["amt_statement"] and row["amt_qbo"],
                  str(row))
            check("so it reads as matched", row["verdict"] == "matched",
                  row["verdict"])
        halves = [r for r in crl["rows"] if r["invoice_number"] == "1"]
        check("no orphaned bare-number duplicate remains", not halves)

    print("\n── penny tolerance ───────────────────────────────────")
    if crl:
        row = next((r for r in crl["rows"]
                    if r["invoice_number"] == "ABC-000006"), None)
        check("a one-cent difference still matches",
              row and row["verdict"] == "matched",
              str(row and row["verdict"]))

    print("\n── pushing bills to QuickBooks ───────────────────────")
    invs = c.get("/api/invoices?pushed=0").get_json()["invoices"]

    unlinked = next((i for i in invs if i["vendor_name"] == "Company STU"), None)
    if unlinked:
        r = c.post(f"/api/invoices/{unlinked['id']}/push", json={})
        d = r.get_json()
        check("push is refused when no QuickBooks vendor is linked",
              not d.get("ok") and "not linked" in (d.get("error") or "").lower(),
              str(d)[:200])

    lowconf = next((i for i in invs if i.get("confidence") == "low"), None)
    if lowconf:
        r = c.post(f"/api/invoices/{lowconf['id']}/push", json={})
        d = r.get_json()
        check("push is refused on a low-confidence parse",
              not d.get("ok") and d.get("needs_override"), str(d)[:200])

    # Pick one QuickBooks does not already hold. Most of the demo invoices are
    # deliberately mirrored in QBO, and pushing those is *supposed* to be
    # refused — choosing one of them would test the duplicate guard twice and
    # never exercise the happy path.
    con = db.get_db()
    try:
        taken = {r["invoice_norm"] for r in
                 con.execute("SELECT invoice_norm FROM qbo_bills").fetchall()}
    finally:
        con.close()

    good = next((i for i in invs
                 if i.get("qbo_vendor_id") and i.get("confidence") != "low"
                 and i.get("amount") and i.get("invoice_norm") not in taken), None)
    check("there is an invoice QuickBooks has never seen", good is not None)
    if good:
        print(f"    pushing {good['invoice_number']} ({good['vendor_name']})")
        r = c.post(f"/api/invoices/{good['id']}/push", json={})
        d = r.get_json()
        check("a clean invoice pushes", d.get("ok"), str(d)[:250])
        r = c.post(f"/api/invoices/{good['id']}/push", json={})
        d = r.get_json()
        check("pushing the same invoice twice is refused",
              not d.get("ok") and "already" in (d.get("error") or "").lower(),
              str(d)[:200])

    print("\n── agents ────────────────────────────────────────────")
    ags = c.get("/api/agents").get_json()["agents"]
    check("built-in agents exist", len(ags) >= 8)

    stmt_agent = next((a for a in ags if a["kind"] == "request_statement"), None)
    if stmt_agent:
        r = c.post(f"/api/agents/{stmt_agent['id']}/run", json={})
        d = r.get_json()
        check("the statement-request agent runs", d.get("ok"), str(d)[:250])
        drafts = c.get("/api/drafts?status=pending").get_json()["drafts"]
        check("it wrote a draft rather than sending", len(drafts) > 0)
        check("the draft explains itself",
              any(x.get("reason") for x in drafts))
        # Running it again must not nag the same vendor twice.
        before = len(drafts)
        c.post(f"/api/agents/{stmt_agent['id']}/run", json={})
        after = len(c.get("/api/drafts?status=pending").get_json()["drafts"])
        check("running it twice does not duplicate the chase", after == before,
              f"{before} → {after}")

    for kind in ("sync_qbo", "run_match", "parse_pending", "request_invoice",
                 "chase_exceptions", "report"):
        a = next((x for x in ags if x["kind"] == kind), None)
        if not a:
            continue
        r = c.post(f"/api/agents/{a['id']}/run", json={})
        d = r.get_json()
        check(f"agent '{kind}' runs", d.get("ok"), str(d)[:200])

    print("\n── drafts ────────────────────────────────────────────")
    drafts = c.get("/api/drafts?status=pending").get_json()["drafts"]
    if drafts:
        did = drafts[0]["id"]
        r = c.post(f"/api/drafts/{did}", json={"action": "update",
                                               "subject": "Edited subject"})
        check("a draft can be edited", r.get_json().get("ok"))
        r = c.post(f"/api/drafts/{did}", json={"action": "send"})
        check("an approved draft sends", r.get_json().get("ok"),
              str(r.get_json())[:200])
        if len(drafts) > 1:
            r = c.post(f"/api/drafts/{drafts[1]['id']}",
                       json={"action": "discard"})
            check("a draft can be discarded", r.get_json().get("ok"))

    print("\n── the workbook ──────────────────────────────────────")
    r = c.post("/api/reports/build", json={})
    d = r.get_json()
    check("the workbook builds", d.get("ok"), str(d)[:250])
    check("it has a tab for every vendor", d.get("vendor_tabs", 0) >= 10,
          f"tabs={d.get('vendor_tabs')}")

    if d.get("ok"):
        from openpyxl import load_workbook
        wb = load_workbook(d["path"])
        names = wb.sheetnames
        check("Summary, Payment plan and Exceptions are present",
              all(n in names for n in ("Summary", "Payment plan", "Exceptions")),
              str(names))
        check("every sheet name is within Excel's 31-char limit",
              all(len(n) <= 31 for n in names),
              str([n for n in names if len(n) > 31]))
        check("sheet names are unique case-insensitively",
              len({n.lower() for n in names}) == len(names))
        pqr = [n for n in names if n.startswith("Company PQR")]
        check("the two 31-char-colliding vendors both got a tab",
              len(pqr) == 2, str(pqr))
        check("a vendor tab carries its rows",
              wb["Company ABC"].max_row > 8)

    rid = c.get("/api/reports").get_json()["reports"][0]["id"]
    r = c.get(f"/api/reports/{rid}/download")
    check("the workbook downloads",
          r.status_code == 200 and len(r.data) > 5000,
          f"status {r.status_code} bytes {len(r.data)}")

    print("\n── settings and connection tests ─────────────────────")
    for name in ("graph", "qbo", "servicetitan", "claude"):
        r = c.post(f"/api/settings/test/{name}")
        check(f"'{name}' reports its connection state",
              r.status_code == 200 and "ok" in (r.get_json() or {}),
              str(r.get_json())[:160])

    r = c.post("/api/settings", json={"settings": {"company_name": "Test Co"}})
    check("a setting saves", r.get_json().get("ok"))
    check("and comes back",
          c.get("/api/settings").get_json()["general"]["company_name"] == "Test Co")

    # A masked secret must never be written back over the real credential.
    db.set_setting("claude_api_key", "sk-real-secret-value", "test")
    masked = next(f["value"] for f in
                  c.get("/api/settings").get_json()["integrations"]["claude"]["fields"]
                  if f["key"] == "claude_api_key")
    check("secrets are masked in the API", masked.startswith("•"), masked)
    c.post("/api/settings", json={"settings": {"claude_api_key": masked}})
    check("saving the mask does not destroy the stored secret",
          db.get_setting("claude_api_key") == "sk-real-secret-value",
          db.get_setting("claude_api_key"))

    print("\n── permissions ───────────────────────────────────────")
    import auth as auth_mod
    auth_mod.create_user("viewer@test.local", "Viewer", "viewonly123", "viewer")
    c2 = app.test_client()
    c2.post("/login", data={"email": "viewer@test.local",
                            "password": "viewonly123"})
    r = c2.post("/api/vendors", json={"name": "Nope"})
    check("a viewer cannot create a vendor", r.status_code == 403)
    r = c2.get("/api/vendors")
    check("but a viewer can read", r.status_code == 200)
    if drafts:
        r = c2.post(f"/api/drafts/{drafts[0]['id']}", json={"action": "send"})
        check("a viewer cannot send email", r.status_code == 403)

    print("\n── vendor guard rails ────────────────────────────────")
    r = c.post("/api/vendors", json={"name": ""})
    check("a nameless vendor is rejected", not r.get_json().get("ok"))
    r = c.post("/api/vendors", json={"name": "Company ABC"})
    check("a duplicate vendor name is rejected",
          not r.get_json().get("ok"), str(r.get_json())[:160])

    vid = c.get("/api/vendors").get_json()["vendors"][0]["id"]
    r = c.delete(f"/api/vendors/{vid}")
    check("deleting a vendor with history asks first",
          r.status_code == 409 and r.get_json().get("needs_force"),
          str(r.get_json())[:200])

    r = c.post("/api/agents", json={"name": "Bad", "kind": "not_a_kind"})
    check("an unknown agent kind is rejected", not r.get_json().get("ok"))
    r = c.post("/api/agents", json={"name": "Bad cron", "kind": "run_match",
                                    "schedule_kind": "cron", "cron": "0 8"})
    check("a malformed cron is rejected and explained",
          not r.get_json().get("ok")
          and "five" in (r.get_json().get("error") or "").lower(),
          str(r.get_json())[:200])

    print("\n" + "═" * 56)
    print(f"  {len(PASS)} passed, {len(FAIL)} failed")
    if FAIL:
        print("\n  Failed:")
        for f in FAIL:
            print(f"    · {f}")
    print("═" * 56 + "\n")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
